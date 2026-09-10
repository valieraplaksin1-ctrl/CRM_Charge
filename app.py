from flask import Flask, render_template, request, jsonify, session, redirect, url_for
from flask_session import Session
from config import BOT_TOKEN, WEB_APP_URL, FLASK_ENV, DEBUG
from database import *
import os
from functools import wraps
from openpyxl import load_workbook

app = Flask(__name__)
app.config['SECRET_KEY'] = 'crm_secret_key_2024'
app.config['SESSION_TYPE'] = 'filesystem'
Session(app)

# Инициализация БД
init_db()

def login_required(f):
    """Декоратор для проверки авторизации"""
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'user_id' not in session:
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return decorated_function

@app.route('/')
def index():
    if 'user_id' in session:
        return redirect(url_for('dashboard'))
    return redirect(url_for('login'))

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        data = request.json
        username = data.get('username')
        password = data.get('password')
        
        user = get_user(username)
        if user and user[2] == password:
            session['user_id'] = user[0]
            session['username'] = user[1]
            return jsonify({'success': True})
        return jsonify({'success': False, 'message': 'Неверные учётные данные'}), 401
    
    return render_template('login.html')

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))

@app.route('/admin', methods=['GET', 'POST'])
def admin():
    if request.method == 'POST':
        data = request.json
        action = data.get('action')
        password = data.get('password')
        
        # Проверка пароля админки
        if password != 'benzolaloh':
            return jsonify({'success': False, 'message': 'Неверный пароль'}), 401
        
        if action == 'get_users':
            users = get_all_users()
            return jsonify({'users': users})
        
        elif action == 'add_user':
            username = data.get('username')
            password = data.get('password')
            if add_user(username, password):
                return jsonify({'success': True})
            return jsonify({'success': False, 'message': 'Пользователь уже существует'}), 400
        
        elif action == 'delete_user':
            user_id = data.get('user_id')
            delete_user(user_id)
            return jsonify({'success': True})
        
        elif action == 'upload_clients':
            # Загрузка клиентов из Excel
            file = request.files.get('file')
            if file:
                try:
                    wb = load_workbook(file)
                    ws = wb.active
                    for row in ws.iter_rows(min_row=2, values_only=True):
                        if row[0]:  # name
                            add_client(row[0], row[1], row[2] if len(row) > 2 else '')
                    return jsonify({'success': True})
                except:
                    return jsonify({'success': False, 'message': 'Ошибка при загрузке файла'}), 400
        
        return jsonify({'success': False}), 400
    
    return render_template('admin.html')

@app.route('/dashboard')
@login_required
def dashboard():
    user_id = session['user_id']
    shift = get_current_shift(user_id)
    shift_active = shift is not None
    duration = get_shift_duration(user_id) if shift_active else 0
    return render_template('dashboard.html', shift_active=shift_active, duration=duration)

@app.route('/api/shift', methods=['POST'])
@login_required
def shift_api():
    user_id = session['user_id']
    action = request.json.get('action')
    
    if action == 'start':
        start_shift(user_id)
        return jsonify({'success': True})
    elif action == 'end':
        end_shift(user_id)
        return jsonify({'success': True})
    elif action == 'status':
        shift = get_current_shift(user_id)
        shift_active = shift is not None
        duration = get_shift_duration(user_id) if shift_active else 0
        return jsonify({'active': shift_active, 'duration': duration})
    
    return jsonify({'success': False}), 400

@app.route('/clients')
@login_required
def clients():
    return render_template('clients.html')

@app.route('/api/clients', methods=['GET', 'POST'])
@login_required
def clients_api():
    user_id = session['user_id']
    
    if request.method == 'GET':
        action = request.args.get('action')
        
        if action == 'next':
            client = get_available_clients()
            if client:
                return jsonify({
                    'id': client[0],
                    'name': client[1],
                    'phone': client[2],
                    'email': client[3]
                })
            return jsonify({'error': 'Нет доступных клиентов'}), 404
        
        elif action == 'working':
            clients = get_working_clients(user_id)
            return jsonify({'clients': [
                {'id': c[0], 'name': c[1], 'phone': c[2], 'email': c[3]} 
                for c in clients
            ]})
    
    elif request.method == 'POST':
        data = request.json
        action = data.get('action')
        client_id = data.get('client_id')
        
        if action == 'skip':
            # Просто получить следующего клиента
            client = get_available_clients()
            if client:
                return jsonify({
                    'id': client[0],
                    'name': client[1],
                    'phone': client[2],
                    'email': client[3]
                })
            return jsonify({'error': 'Нет доступных клиентов'}), 404
        
        elif action == 'add_comment':
            comment = data.get('comment')
            add_comment(client_id, user_id, comment)
            return jsonify({'success': True})
        
        elif action == 'take':
            assign_client(client_id, user_id)
            return jsonify({'success': True})
        
        elif action == 'delete':
            delete_client(client_id)
            return jsonify({'success': True})
    
    return jsonify({'success': False}), 400

@app.route('/working-clients')
@login_required
def working_clients():
    return render_template('working_clients.html')

@app.route('/api/working-clients/<int:client_id>', methods=['GET', 'POST'])
@login_required
def working_client_api(client_id):
    user_id = session['user_id']
    
    if request.method == 'GET':
        comments = get_comments(client_id)
        return jsonify({'comments': [
            {'user': get_username(c[0]), 'text': c[1], 'date': c[2]}
            for c in comments
        ]})
    
    elif request.method == 'POST':
        data = request.json
        action = data.get('action')
        
        if action == 'add_comment':
            comment = data.get('comment')
            add_comment(client_id, user_id, comment)
            return jsonify({'success': True})
        
        elif action == 'delete':
            delete_client(client_id)
            return jsonify({'success': True})
    
    return jsonify({'success': False}), 400

if __name__ == '__main__':
    app.run(debug=DEBUG, host='0.0.0.0', port=5000)
