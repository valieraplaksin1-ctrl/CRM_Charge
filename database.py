import sqlite3
from datetime import datetime
import os
import random
from openpyxl import load_workbook

DB_PATH = 'crm.db'
CLIENTS_FILE = 'data/clients.xlsx'

def init_db():
    """Инициализация базы данных"""
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    
    # Таблица пользователей
    c.execute('''
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY,
            username TEXT UNIQUE NOT NULL,
            password TEXT NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    ''')
    
    # Таблица рабочих смен
    c.execute('''
        CREATE TABLE IF NOT EXISTS shifts (
            id INTEGER PRIMARY KEY,
            user_id INTEGER NOT NULL,
            start_time TIMESTAMP,
            end_time TIMESTAMP,
            FOREIGN KEY(user_id) REFERENCES users(id)
        )
    ''')
    
    # Таблица клиентов
    c.execute('''
        CREATE TABLE IF NOT EXISTS clients (
            id INTEGER PRIMARY KEY,
            name TEXT NOT NULL,
            phone TEXT NOT NULL,
            email TEXT,
            status TEXT DEFAULT 'available',
            assigned_to INTEGER,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY(assigned_to) REFERENCES users(id)
        )
    ''')
    
    # Таблица комментариев
    c.execute('''
        CREATE TABLE IF NOT EXISTS comments (
            id INTEGER PRIMARY KEY,
            client_id INTEGER NOT NULL,
            phone TEXT,
            user_id INTEGER NOT NULL,
            comment TEXT NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY(client_id) REFERENCES clients(id),
            FOREIGN KEY(user_id) REFERENCES users(id)
        )
    ''')
    
    # Таблица текущего клиента агента
    c.execute('''
        CREATE TABLE IF NOT EXISTS current_clients (
            id INTEGER PRIMARY KEY,
            user_id INTEGER UNIQUE NOT NULL,
            client_id INTEGER NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY(user_id) REFERENCES users(id),
            FOREIGN KEY(client_id) REFERENCES clients(id)
        )
    ''')
    
    conn.commit()
    conn.close()
    
    # Загружаем клиентов из файла если его нет в БД
    load_clients_from_file()

def load_clients_from_file():
    """Загружает клиентов из Excel файла в папке data/"""
    if not os.path.exists(CLIENTS_FILE):
        return
    
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    
    # Проверяем есть ли уже клиенты
    c.execute('SELECT COUNT(*) FROM clients')
    if c.fetchone()[0] > 0:
        conn.close()
        return
    
    try:
        wb = load_workbook(CLIENTS_FILE)
        ws = wb.active
        
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0]:  # Если есть имя
                name = str(row[0]).strip()
                phone = str(row[1]).strip() if row[1] else ''
                email = str(row[2]).strip() if row[2] else ''
                
                c.execute(
                    'INSERT INTO clients (name, phone, email, status) VALUES (?, ?, ?, ?)',
                    (name, phone, email, 'available')
                )
        
        conn.commit()
        print(f"✅ Загружено клиентов из {CLIENTS_FILE}")
    except Exception as e:
        print(f"❌ Ошибка при загрузке клиентов: {e}")
    finally:
        conn.close()

def get_user(username):
    """Получить пользователя по имени"""
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    c.execute('SELECT id, username, password FROM users WHERE username = ?', (username,))
    result = c.fetchone()
    conn.close()
    return result

def get_user_by_id(user_id):
    """Получить пользователя по ID"""
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    c.execute('SELECT id, username, password FROM users WHERE id = ?', (user_id,))
    result = c.fetchone()
    conn.close()
    return result

def add_user(username, password):
    """Добавить нового пользователя"""
    try:
        conn = sqlite3.connect(DB_PATH)
        c = conn.cursor()
        c.execute('INSERT INTO users (username, password) VALUES (?, ?)', (username, password))
        conn.commit()
        conn.close()
        return True
    except:
        return False

def update_user_password(user_id, new_password):
    """Обновить пароль пользователя"""
    try:
        conn = sqlite3.connect(DB_PATH)
        c = conn.cursor()
        c.execute('UPDATE users SET password = ? WHERE id = ?', (new_password, user_id))
        conn.commit()
        conn.close()
        return True
    except:
        return False

def delete_user(user_id):
    """Удалить пользователя"""
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    c.execute('DELETE FROM users WHERE id = ?', (user_id,))
    conn.commit()
    conn.close()

def get_all_users():
    """Получить всех пользователей"""
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    c.execute('SELECT id, username, password, created_at FROM users')
    results = c.fetchall()
    conn.close()
    return results

def start_shift(user_id):
    """Начать рабочую смену"""
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    c.execute('INSERT INTO shifts (user_id, start_time) VALUES (?, ?)', (user_id, datetime.now()))
    conn.commit()
    shift_id = c.lastrowid
    conn.close()
    return shift_id

def end_shift(user_id):
    """Закончить рабочую смену"""
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    c.execute('UPDATE shifts SET end_time = ? WHERE user_id = ? AND end_time IS NULL', (datetime.now(), user_id))
    conn.commit()
    conn.close()

def get_current_shift(user_id):
    """Получить текущую смену"""
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    c.execute('SELECT id, start_time FROM shifts WHERE user_id = ? AND end_time IS NULL', (user_id,))
    result = c.fetchone()
    conn.close()
    return result

def get_shift_duration(user_id):
    """Получить длительность текущей смены в секундах"""
    shift = get_current_shift(user_id)
    if not shift:
        return 0
    start_time = datetime.fromisoformat(shift[1])
    duration = (datetime.now() - start_time).total_seconds()
    return int(duration)

def get_today_shift(user_id):
    """Получить смену за сегодня"""
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    c.execute('''
        SELECT SUM(CAST((julianday(end_time) - julianday(start_time)) * 86400 AS INTEGER))
        FROM shifts
        WHERE user_id = ? AND DATE(start_time) = DATE('now')
    ''', (user_id,))
    result = c.fetchone()
    conn.close()
    return result[0] if result[0] else 0

def add_client(name, phone, email):
    """Добавить клиента"""
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    c.execute('INSERT INTO clients (name, phone, email, status) VALUES (?, ?, ?, ?)', 
              (name, phone, email, 'available'))
    conn.commit()
    client_id = c.lastrowid
    conn.close()
    return client_id

def get_current_client(user_id):
    """Получить текущего клиента агента"""
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    c.execute('''
        SELECT c.id, c.name, c.phone, c.email 
        FROM current_clients cc
        JOIN clients c ON cc.client_id = c.id
        WHERE cc.user_id = ?
    ''', (user_id,))
    result = c.fetchone()
    conn.close()
    return result

def set_current_client(user_id, client_id):
    """Сохранить текущего клиента агента"""
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    c.execute('INSERT OR REPLACE INTO current_clients (user_id, client_id) VALUES (?, ?)', 
              (user_id, client_id))
    conn.commit()
    conn.close()

def get_available_clients():
    """Получить случайного доступного клиента"""
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    
    # Получаем всех доступных клиентов
    c.execute('SELECT id, name, phone, email FROM clients WHERE status = "available"')
    clients = c.fetchall()
    conn.close()
    
    # Возвращаем случайного клиента
    if clients:
        return random.choice(clients)
    return None

def get_client_by_id(client_id):
    """Получить данные клиента по ID"""
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    c.execute('SELECT id, name, phone, email FROM clients WHERE id = ?', (client_id,))
    result = c.fetchone()
    conn.close()
    return result

def get_working_clients(user_id):
    """Получить клиентов в работе у пользователя"""
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    c.execute('SELECT id, name, phone, email FROM clients WHERE status = "working" AND assigned_to = ?', (user_id,))
    results = c.fetchall()
    conn.close()
    return results

def assign_client(client_id, user_id):
    """Назначить клиента пользователю"""
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    c.execute('UPDATE clients SET status = ?, assigned_to = ? WHERE id = ?', 
              ('working', user_id, client_id))
    conn.commit()
    conn.close()

def delete_client(client_id):
    """Удалить клиента"""
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    c.execute('DELETE FROM clients WHERE id = ?', (client_id,))
    conn.commit()
    conn.close()

def add_comment(client_id, user_id, comment):
    """Добавить комментарий"""
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    
    # Получаем номер телефона клиента для привязки
    c.execute('SELECT phone FROM clients WHERE id = ?', (client_id,))
    result = c.fetchone()
    phone = result[0] if result else ''
    
    c.execute('INSERT INTO comments (client_id, phone, user_id, comment) VALUES (?, ?, ?, ?)', 
              (client_id, phone, user_id, comment))
    conn.commit()
    conn.close()

def get_comments(client_id):
    """Получить комментарии к клиенту"""
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    c.execute('SELECT user_id, comment, created_at FROM comments WHERE client_id = ? ORDER BY created_at DESC', (client_id,))
    results = c.fetchall()
    conn.close()
    return results

def get_all_comments_by_phone(phone):
    """Получить все комментарии по номеру телефона (для истории)"""
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    c.execute('SELECT user_id, comment, created_at FROM comments WHERE phone = ? ORDER BY created_at DESC', (phone,))
    results = c.fetchall()
    conn.close()
    return results

def get_last_comment(client_id):
    """Получить последний комментарий к клиенту с датой"""
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    c.execute('SELECT user_id, comment, created_at FROM comments WHERE client_id = ? ORDER BY created_at DESC LIMIT 1', (client_id,))
    result = c.fetchone()
    conn.close()
    return result

def get_username(user_id):
    """Получить имя пользователя по ID"""
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    c.execute('SELECT username FROM users WHERE id = ?', (user_id,))
    result = c.fetchone()
    conn.close()
    return result[0] if result else 'Unknown'

def count_user_calls_today(user_id):
    """Считать количество клиентов прозвонено сегодня"""
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    c.execute('''
        SELECT COUNT(*) FROM comments 
        WHERE user_id = ? AND DATE(created_at) = DATE('now')
    ''', (user_id,))
    result = c.fetchone()
    conn.close()
    return result[0] if result else 0
